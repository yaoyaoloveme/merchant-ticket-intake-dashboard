#!/usr/bin/env python3
"""客诉稽查分析：只读同步、事实统计、历史版本和内网访问控制。"""

from __future__ import annotations

import json
import math
import os
import re
import secrets
import sqlite3
import threading
import time
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from html import unescape
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin

import requests
from flask import Blueprint, Response, jsonify, make_response, redirect, render_template, request, url_for
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import load_workbook
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from werkzeug.security import check_password_hash


ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
DB_PATH = DATA_DIR / "tickets.db"
DEFAULT_SOURCE_BASE = "http://172.25.11.31:8080/ticket-dashboard/"
SOURCE_BASE = os.environ.get("COMPLAINT_AUDIT_SOURCE_URL", DEFAULT_SOURCE_BASE).rstrip("/") + "/"
BACKFILL_PATH = "api/complaint-audits/backfill"
CHANGES_PATH = "api/complaint-audits/changes"
SYNC_INTERVAL_SECONDS = 60
OVERLAP_MINUTES = 5
SESSION_COOKIE = "complaint_audit_session"
SESSION_MAX_AGE_SECONDS = 12 * 60 * 60
LOGIN_WINDOW_MINUTES = 15
LOGIN_FAILURE_LIMIT = 5
SYNC_KEY = "forge-complaint-audits"

bp = Blueprint("complaint_audit", __name__)
SYNC_LOCK = threading.Lock()
AUTO_SYNC_STARTED = False
AUTO_STATE: dict[str, Any] = {
    "running": False,
    "message": "尚未同步",
    "last_started_at": None,
    "last_finished_at": None,
    "last_error": None,
    "next_run_at": None,
}

RESULT_YES = {"是", "true", "1", "yes", "y", "成立", "客诉成立"}
RESULT_NO = {"否", "false", "0", "no", "n", "不成立", "客诉不成立"}
APPROVED_STATUSES = {"审核通过", "已审核", "审核完成", "通过"}
PENDING_STATUSES = {"待审核", "审核中", "待复核"}
RETURNED_STATUSES = {"审核退回", "审核不通过", "已退回", "退回", "审核驳回", "已驳回"}
RESPONSIBILITY_SUBJECTS = ("人工客服责任", "AI工作台责任", "系统/平台责任", "其他责任", "责任待确认")
RESPONSIBILITY_ALIASES = {
    "人工客服责任": "人工客服责任",
    "人工客服": "人工客服责任",
    "客服责任": "人工客服责任",
    "值守客服责任": "人工客服责任",
    "ai工作台责任": "AI工作台责任",
    "ai工作台": "AI工作台责任",
    "ai责任": "AI工作台责任",
    "系统/平台责任": "系统/平台责任",
    "系统责任": "系统/平台责任",
    "平台责任": "系统/平台责任",
    "其他责任": "其他责任",
    "其他": "其他责任",
}

EXCEL_REQUIRED_COLUMNS = {
    "反馈来源", "店铺ID", "工单ID", "稽查单ID", "反馈时间", "昵称", "客诉分类别名",
    "原因分类别名", "客诉原因描述", "稽查状态", "完成处理时间", "审核状态",
    "完成审核时间", "最后修改日期", "ID",
}
EXCEL_RESPONSIBILITY_BRANCH_MAP = {
    "值守客服原因": "人工客服责任",
    "AI工作台值守": "AI工作台责任",
    "AI值守原因": "AI工作台责任",
    "系统原因": "系统/平台责任",
    "商家原因": "其他责任",
    "顾客原因": "其他责任",
    "服务管家原因": "其他责任",
    "虚拟值守": "其他责任",
    "核实结果无异常": "其他责任",
    "重复反馈": "其他责任",
}
EXCEL_TIME_BASIS = "完成处理时间（临时替代稽查完成时间）"
MAX_IMPORT_BYTES = 25 * 1024 * 1024

DESCRIPTION_THEME_RULES = (
    ("描述未填写或过于笼统", "描述缺失或过于笼统", re.compile(
        r"^(?:|未填写|无|暂无|这单有问题|特殊诉求|商家特殊需求|商家特殊要求)$", re.I)),
    ("重复反馈或已有稽查结果", "包含明确核实结果或状态", re.compile(
        r"重复(?:反馈|工单|提单|发起|稽查单)?|多渠道重复|已有稽查结果|已经稽查过|反馈前已稽查|已稽查完成|打标已稽查|已重新发送稽查单|商家反馈前已被稽查|原稽查单", re.I)),
    ("商家确认无异常或撤回终止", "包含明确核实结果或状态", re.compile(
        r"商家(?:又)?(?:反馈|表示|告知).{0,8}(?:无异常|没问题|看错|已处理|可以关单|正常)|(?:老板|店主|管家).{0,8}(?:无异常|没问题|看错)|商家.{0,6}(?:要求|反馈)?终止|(?:要求|协助|辛苦)?终止(?:定损|追缴单|稽查单|工单)?|要求关单|让关单|无需(?:终止|追讨|赔付|核实)|不再(?:追讨|赔付)|已终止", re.I)),
    ("订单支付或商品数量核实正常", "包含明确核实结果或状态", re.compile(
        r"支付成功|付款成功|正常(?:付款|支付|加购|选购)|已经付款|已付款|已支付|补付|现金支付|转账付款|二次进店付款|(?<!未)支付离店|扫码支付|并未漏付|没有漏付|无漏付|未漏付|未少付|顾客无漏付|订单(?:均)?无异常|核实订单无异常|数量对|无差价|金额一致|商品正常|香烟一致|商品.{0,8}放回|已放回|未带走|自带商品|自己.{0,6}带进|等价置换.{0,6}无问题", re.I)),
    ("反馈时间、提单或核实对象不匹配", "包含明确核实结果或状态", re.compile(
        r"时间.{0,6}(?:有误|错误|不准确|不符)|店内.{0,5}(?:无人|无顾客|没有顾客)|反馈时间.{0,8}(?:无人|无顾客|没有顾客)|时间点.{0,8}(?:无人|无顾客|没有顾客)|无顾客|没有顾客|无人进入|没有人进入|未核实顾客进店|漏付顾客为|后面进店|待商家核实时间|还没提供时间|非本店稽查单|提单(?:信息)?(?:错误|有误|存在偏差)|订单号错误|非值班店长问题", re.I)),
    ("视频或监控条件不足无法核实", "包含明确核实结果或状态", re.compile(
        r"黑屏|滚动覆盖|视频.{0,10}(?:无法|不能|缺失|问题|无视频)|监控.{0,10}(?:无法|不能|缺失|问题|死角)|无视频|视频已.{0,4}覆盖|死角|遮挡.{0,8}无法核实|无法核实", re.I)),
    ("非值守或相关功能未开启", "包含明确核实结果或状态", re.compile(
        r"未开(?:启)?值守|没开.{0,8}值守|关闭.{0,8}值守|非值守|不属于值守问题|未开启定损|未定损|商家在店.{0,8}值守|店主在店", re.I)),
    ("客服或值守已有介入", "包含明确核实结果或状态", re.compile(
        r"正常介入|客服.{0,10}(?:有介入|已介入|正常介入|正常引导|进行介入|已响应|现场拦截|正常值守|有开麦|已开麦|有引导|已引导)|AI客服.{0,8}介入|有介入|已介入|有开麦|有引导", re.I)),
    ("核实未发现异常", "包含明确核实结果或状态", re.compile(
        r"无异常|未发现(?:其他)?异常|未见异常|未查询到异常|未核实到.{0,10}异常|核实.{0,8}无误|批量.{0,10}(?:无异常|未查询到|未见)|全量.{0,10}(?:无异常|未查询到|未见)|暂未发现|未见(?:少付|漏付)|未发现.{0,8}(?:漏付|少付|偷盗|丢失)|未查询到.{0,8}(?:漏付|少付|偷盗|丢失)", re.I)),
    ("仅描述核实诉求，未写明结论", "仅描述诉求或处理动作", re.compile(
        r"商家要求核实|要求核实|稽查优先核实|资损问题|资损情况|资损进度|咨询订单|老板反馈(?:漏付|商品丢失|香烟丢失)|商家反馈(?:漏付|香烟丢失|商品丢失)|反馈丢失|商家说烟丢|询问具体情况|核实具体", re.I)),
    ("仅描述视频诉求，未写明结论", "仅描述诉求或处理动作", re.compile(
        r"^视频$|调取视频|提供视频|发送视频|发视频|需要视频|视频调取|视频下载|下载视频|调取监控|商家需求调取视频", re.I)),
    ("仅描述赔付或定损诉求，未写明结论", "仅描述诉求或处理动作", re.compile(
        r"赔付|定损|追讨|差价", re.I)),
    ("仅描述介入诉求，未写明结论", "仅描述诉求或处理动作", re.compile(
        r"介入问题|介入不及时|未及时介入|询问未介入|询问未及时介入", re.I)),
    ("描述仍含异常或待处理事项", "描述仍含异常或待处理", re.compile(
        r"漏付|少付|偷盗|丢失|未支付|有异常|卡顿|拿错|损坏|多付|未介入|未及时|超时|未稽查|不认可|已成立", re.I)),
    ("仅记录通知或一般处理动作", "仅描述诉求或处理动作", re.compile(
        r"发送(?:至)?(?:老板|商家群|群内)|已?优先发送|告知商家|优先(?:告知|反馈|处理)|客服优先反馈|已留言|已处理|处理完成|拉黑|二维码|已驳回|已补充|扫描|已联系|已致电", re.I)),
)

RECORD_FIELDS = (
    "source_record_id", "inspect_id", "ticket_id", "shop_id", "shop_name", "brand",
    "inspect_status", "inspect_completed_at", "complaint_result", "review_status",
    "complained_agent", "reason_node_id", "reason_path", "responsibility_subject",
    "reason_description", "inspect_result", "abnormal_state", "abnormal_type",
    "handled_at", "reviewed_at", "source_updated_at", "deleted",
)
DISPLAY_NAMES = {
    "source_record_id": "投诉记录ID", "inspect_id": "稽查单ID", "ticket_id": "工单ID",
    "shop_id": "店铺ID", "shop_name": "店铺", "brand": "品牌", "inspect_status": "稽查状态",
    "inspect_completed_at": "稽查完成时间", "complaint_result": "客诉结果",
    "review_status": "审核状态", "complained_agent": "被投诉客服",
    "reason_node_id": "原因节点ID", "reason_path": "完整责任分类路径",
    "responsibility_subject": "责任主体", "reason_description": "原因描述",
    "inspect_result": "稽查结论", "abnormal_state": "是否异常", "abnormal_type": "异常类型",
    "handled_at": "处理时间", "reviewed_at": "审核时间", "source_updated_at": "源数据更新时间",
    "deleted": "删除状态",
}


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def strip_html(value: Any) -> str:
    text = unescape("" if value is None else str(value))
    text = re.sub(r"<\s*br\s*/?\s*>", "\n", text, flags=re.I)
    text = re.sub(r"</\s*p\s*>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    return clean(text)


def split_reason_path(value: Any) -> list[str]:
    text = clean(value)
    if not text:
        return []
    if "##" in text:
        return [clean(part) for part in text.split("##") if clean(part)]
    return [clean(part) for part in re.split(r"\s+/\s+", text) if clean(part)]


def normalized_reason_path(value: Any) -> str:
    return " / ".join(split_reason_path(value))


def normalized_description(value: Any) -> str:
    return clean(value).replace("\ufeff", "").strip(" ,，。；;")


def description_theme_for(value: Any) -> tuple[str, str]:
    text = normalized_description(value)
    for theme, group, pattern in DESCRIPTION_THEME_RULES:
        if pattern.search(text):
            return theme, group
    return "其他未归类原因描述", "其他未归类"


def source_mode() -> str:
    return "excel" if os.environ.get("COMPLAINT_AUDIT_SOURCE_MODE", "api").strip().lower() == "excel" else "api"


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def parse_datetime(value: Any) -> datetime | None:
    text = clean(value)
    if not text:
        return None
    text = text.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        for fmt in ("%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
    return None


def normalize_datetime(value: Any) -> str:
    parsed = parse_datetime(value)
    if parsed:
        return parsed.isoformat(timespec="seconds")
    return clean(value)


def subtract_minutes(value: str, minutes: int) -> str:
    parsed = parse_datetime(value)
    return (parsed - timedelta(minutes=minutes)).isoformat(timespec="seconds") if parsed else value


def month_start(month: str) -> date:
    if not re.fullmatch(r"\d{4}-\d{2}", month):
        raise ValueError("月份格式应为 YYYY-MM")
    return date.fromisoformat(month + "-01")


def shift_month(day: date, delta: int) -> date:
    index = day.year * 12 + day.month - 1 + delta
    return date(index // 12, index % 12 + 1, 1)


def month_end(day: date) -> date:
    return shift_month(day, 1) - timedelta(days=1)


def six_month_start(today: date | None = None) -> date:
    return shift_month((today or date.today()).replace(day=1), -5)


def get_db() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def init_db() -> None:
    conn = get_db()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS complaint_audit_records (
          source_record_id TEXT PRIMARY KEY,
          inspect_id TEXT NOT NULL,
          ticket_id TEXT,
          shop_id TEXT,
          shop_name TEXT,
          brand TEXT,
          inspect_status TEXT,
          inspect_completed_at TEXT,
          complaint_result TEXT,
          review_status TEXT,
          complained_agent TEXT,
          reason_node_id TEXT,
          reason_path TEXT,
          responsibility_subject TEXT,
          reason_description TEXT,
          inspect_result TEXT,
          abnormal_state TEXT,
          abnormal_type TEXT,
          handled_at TEXT,
          reviewed_at TEXT,
          source_updated_at TEXT,
          deleted INTEGER NOT NULL DEFAULT 0,
          source_payload TEXT NOT NULL,
          local_updated_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_complaint_inspect_id ON complaint_audit_records(inspect_id);
        CREATE INDEX IF NOT EXISTS idx_complaint_completed_at ON complaint_audit_records(inspect_completed_at);
        CREATE INDEX IF NOT EXISTS idx_complaint_source_updated ON complaint_audit_records(source_updated_at);
        CREATE INDEX IF NOT EXISTS idx_complaint_agent ON complaint_audit_records(complained_agent);

        CREATE TABLE IF NOT EXISTS complaint_audit_versions (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          source_record_id TEXT NOT NULL,
          inspect_id TEXT NOT NULL,
          version_no INTEGER NOT NULL,
          changed_at TEXT NOT NULL,
          changed_fields TEXT NOT NULL,
          before_json TEXT,
          after_json TEXT NOT NULL
        );
        CREATE UNIQUE INDEX IF NOT EXISTS idx_complaint_version_unique
          ON complaint_audit_versions(source_record_id, version_no);
        CREATE INDEX IF NOT EXISTS idx_complaint_version_inspect
          ON complaint_audit_versions(inspect_id, changed_at);

        CREATE TABLE IF NOT EXISTS complaint_audit_sync_cursor (
          source_key TEXT PRIMARY KEY,
          cursor TEXT,
          last_source_updated_at TEXT,
          last_success_at TEXT,
          last_reconcile_date TEXT,
          data_version INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS complaint_audit_sync_runs (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          mode TEXT NOT NULL,
          started_at TEXT NOT NULL,
          finished_at TEXT,
          status TEXT NOT NULL,
          received_count INTEGER NOT NULL DEFAULT 0,
          changed_count INTEGER NOT NULL DEFAULT 0,
          deleted_count INTEGER NOT NULL DEFAULT 0,
          duration_ms INTEGER,
          message TEXT
        );

        CREATE TABLE IF NOT EXISTS complaint_audit_access_log (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          attempted_at TEXT NOT NULL,
          client_addr TEXT NOT NULL,
          result TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_complaint_access_addr_time
          ON complaint_audit_access_log(client_addr, attempted_at);
        """
    )
    conn.execute(
        "INSERT OR IGNORE INTO complaint_audit_sync_cursor "
        "(source_key,cursor,last_source_updated_at,last_success_at,last_reconcile_date,data_version) "
        "VALUES (?,?,?,?,?,0)",
        (SYNC_KEY, "", "", "", ""),
    )
    conn.commit()
    conn.close()


def pick(record: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in record and record[key] is not None:
            return record[key]
    return None


def normalize_result(value: Any) -> str:
    if value is True:
        return "是"
    if value is False:
        return "否"
    text = clean(value).lower()
    if text in RESULT_YES:
        return "是"
    if text in RESULT_NO:
        return "否"
    return ""


def normalize_review_status(value: Any) -> str:
    text = clean(value)
    lowered = text.lower()
    if text in APPROVED_STATUSES or lowered in {item.lower() for item in APPROVED_STATUSES}:
        return "审核通过"
    if text in PENDING_STATUSES or lowered in {item.lower() for item in PENDING_STATUSES}:
        return "待审核"
    if text in RETURNED_STATUSES or lowered in {item.lower() for item in RETURNED_STATUSES}:
        return "审核退回"
    return text


def responsibility_node_map() -> dict[str, str]:
    raw = os.environ.get("COMPLAINT_AUDIT_RESPONSIBILITY_NODE_MAP", "").strip()
    if not raw:
        return {}
    try:
        if raw.startswith("{"):
            loaded = json.loads(raw)
        else:
            loaded = json.loads(Path(raw).expanduser().read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return {clean(key): canonical_subject(value) for key, value in loaded.items() if clean(key)}


def canonical_subject(value: Any) -> str:
    text = clean(value)
    if not text:
        return ""
    return RESPONSIBILITY_ALIASES.get(text.lower(), RESPONSIBILITY_ALIASES.get(text, ""))


def resolve_subject(explicit: Any, node_id: Any) -> str:
    subject = canonical_subject(explicit)
    if subject:
        return subject
    return responsibility_node_map().get(clean(node_id), "")


def normalize_record(raw: dict[str, Any]) -> dict[str, Any]:
    record_id = clean(pick(raw, "complaint_record_id", "complaintRecordId", "source_record_id", "record_id", "recordId", "id"))
    inspect_id = clean(pick(raw, "inspect_id", "inspectId", "inspection_id", "inspectionId", "inspection_order_id", "inspectionOrderId"))
    if not record_id:
        raise ValueError("上游记录缺少投诉记录ID")
    deleted_value = pick(raw, "deleted", "is_deleted", "isDeleted")
    deleted = int(deleted_value is True or clean(deleted_value).lower() in {"1", "true", "yes", "是", "已删除"})
    reason_node_id = clean(pick(raw, "reason_node_id", "reasonNodeId", "category_node_id", "categoryNodeId"))
    explicit_subject = pick(raw, "responsibility_subject", "responsibilitySubject", "responsibility_type", "responsibilityType")
    normalized = {
        "source_record_id": record_id,
        "inspect_id": inspect_id,
        "ticket_id": clean(pick(raw, "ticket_id", "ticketId", "work_order_id", "workOrderId")),
        "shop_id": clean(pick(raw, "shop_id", "shopId", "store_id", "storeId")),
        "shop_name": clean(pick(raw, "shop_name", "shopName", "store_name", "storeName")),
        "brand": clean(pick(raw, "brand", "brand_name", "brandName")),
        "inspect_status": clean(pick(raw, "inspect_status", "inspectStatus", "inspection_status", "inspectionStatus")),
        "inspect_completed_at": normalize_datetime(pick(raw, "inspect_completed_at", "inspectCompletedAt", "inspection_completed_at", "inspectionCompletedAt", "completed_at", "completedAt")),
        "complaint_result": normalize_result(pick(raw, "complaint_result", "complaintResult", "is_complaint_true", "isComplaintTrue")),
        "review_status": normalize_review_status(pick(raw, "review_status", "reviewStatus", "audit_status", "auditStatus")),
        "complained_agent": clean(pick(raw, "complained_agent", "complainedAgent", "complained_customer_service", "complainedCustomerService", "customer_service_name", "customerServiceName")),
        "reason_node_id": reason_node_id,
        "reason_path": clean(pick(raw, "reason_path", "reasonPath", "category_path", "categoryPath", "full_category_path", "fullCategoryPath")),
        "responsibility_subject": resolve_subject(explicit_subject, reason_node_id),
        "reason_description": clean(pick(raw, "reason_description", "reasonDescription", "complaint_reason", "complaintReason")),
        "inspect_result": clean(pick(raw, "inspect_result", "inspectResult", "inspection_result", "inspectionResult", "conclusion")),
        "abnormal_state": clean(pick(raw, "abnormal_state", "abnormalState", "is_abnormal", "isAbnormal")),
        "abnormal_type": clean(pick(raw, "abnormal_type", "abnormalType")),
        "handled_at": normalize_datetime(pick(raw, "handled_at", "handledAt", "processed_at", "processedAt")),
        "reviewed_at": normalize_datetime(pick(raw, "reviewed_at", "reviewedAt", "review_completed_at", "reviewCompletedAt")),
        "source_updated_at": normalize_datetime(pick(raw, "source_updated_at", "sourceUpdatedAt", "updated_at", "updatedAt", "last_modified_at", "lastModifiedAt")),
        "deleted": deleted,
        "source_payload": json.dumps(raw, ensure_ascii=False, sort_keys=True, default=str),
        "local_updated_at": now_iso(),
    }
    if not inspect_id and not deleted:
        raise ValueError(f"投诉记录 {record_id} 缺少稽查单ID")
    return normalized


def record_snapshot(record: dict[str, Any] | sqlite3.Row) -> dict[str, Any]:
    return {field: record[field] for field in RECORD_FIELDS}


def upsert_record(conn: sqlite3.Connection, record: dict[str, Any]) -> tuple[bool, bool]:
    old_row = conn.execute(
        "SELECT * FROM complaint_audit_records WHERE source_record_id=?", (record["source_record_id"],)
    ).fetchone()
    if old_row and record["deleted"] and not record["inspect_id"]:
        for field in RECORD_FIELDS:
            if field not in {"deleted", "source_updated_at"} and not record.get(field):
                record[field] = old_row[field]
    before = record_snapshot(old_row) if old_row else None
    after = record_snapshot(record)
    changed_fields = [field for field in RECORD_FIELDS if before is None or before.get(field) != after.get(field)]
    if not changed_fields:
        return False, bool(record["deleted"])
    if old_row:
        assignments = ",".join(f"{field}=?" for field in RECORD_FIELDS[1:])
        conn.execute(
            f"UPDATE complaint_audit_records SET {assignments},source_payload=?,local_updated_at=? WHERE source_record_id=?",
            tuple(record[field] for field in RECORD_FIELDS[1:])
            + (record["source_payload"], record["local_updated_at"], record["source_record_id"]),
        )
        version_no = conn.execute(
            "SELECT COALESCE(MAX(version_no),0)+1 FROM complaint_audit_versions WHERE source_record_id=?",
            (record["source_record_id"],),
        ).fetchone()[0]
    else:
        columns = ",".join(RECORD_FIELDS) + ",source_payload,local_updated_at"
        placeholders = ",".join("?" for _ in range(len(RECORD_FIELDS) + 2))
        conn.execute(
            f"INSERT INTO complaint_audit_records ({columns}) VALUES ({placeholders})",
            tuple(record[field] for field in RECORD_FIELDS) + (record["source_payload"], record["local_updated_at"]),
        )
        version_no = 1
    conn.execute(
        "INSERT INTO complaint_audit_versions "
        "(source_record_id,inspect_id,version_no,changed_at,changed_fields,before_json,after_json) "
        "VALUES (?,?,?,?,?,?,?)",
        (
            record["source_record_id"], record["inspect_id"], version_no, now_iso(),
            json.dumps(changed_fields, ensure_ascii=False),
            json.dumps(before, ensure_ascii=False) if before else None,
            json.dumps(after, ensure_ascii=False),
        ),
    )
    return True, bool(record["deleted"])


def mark_deleted(conn: sqlite3.Connection, source_record_id: str, source_updated_at: str = "") -> tuple[bool, bool]:
    old = conn.execute(
        "SELECT * FROM complaint_audit_records WHERE source_record_id=?", (clean(source_record_id),)
    ).fetchone()
    if not old or old["deleted"]:
        return False, False
    record = dict(old)
    record.update(
        {
            "deleted": 1,
            "source_updated_at": normalize_datetime(source_updated_at) or old["source_updated_at"],
            "source_payload": old["source_payload"],
            "local_updated_at": now_iso(),
        }
    )
    return upsert_record(conn, record)


def source_session() -> requests.Session:
    session = requests.Session()
    retry = Retry(
        total=3, connect=3, read=3, backoff_factor=0.7,
        status_forcelist=(408, 429, 500, 502, 503, 504), allowed_methods=("GET",),
    )
    session.mount("http://", HTTPAdapter(max_retries=retry))
    session.mount("https://", HTTPAdapter(max_retries=retry))
    session.headers.update({"Accept": "application/json", "User-Agent": "complaint-audit-dashboard/1.0"})
    return session


def unpack_page(payload: Any) -> tuple[list[dict[str, Any]], list[Any], str, bool, str]:
    root = payload if isinstance(payload, dict) else {}
    data = root.get("data", root)
    if isinstance(data, list):
        records, meta = data, root
    elif isinstance(data, dict):
        records = next((data[key] for key in ("records", "items", "results", "list") if isinstance(data.get(key), list)), [])
        meta = {**root, **data}
    else:
        records, meta = [], root
    deleted = next((meta[key] for key in ("deleted_ids", "deletedIds", "deletions") if isinstance(meta.get(key), list)), [])
    next_cursor = clean(pick(meta, "next_cursor", "nextCursor", "page_cursor", "pageCursor"))
    sync_cursor = clean(pick(meta, "sync_cursor", "syncCursor", "watermark_cursor", "watermarkCursor", "cursor"))
    has_more_raw = pick(meta, "has_more", "hasMore")
    has_more = has_more_raw is True or clean(has_more_raw).lower() in {"1", "true", "yes"}
    return [item for item in records if isinstance(item, dict)], deleted, next_cursor, has_more, sync_cursor


def fetch_pages(path: str, params: dict[str, Any]) -> tuple[list[dict[str, Any]], list[Any], str]:
    session = source_session()
    all_records: list[dict[str, Any]] = []
    all_deleted: list[Any] = []
    page_cursor = ""
    final_sync_cursor = ""
    page = 1
    seen_cursors: set[str] = set()
    for _ in range(10000):
        query = {**params, "page_size": 500, "page": page}
        if page_cursor:
            query["page_cursor"] = page_cursor
        response = session.get(urljoin(SOURCE_BASE, path), params=query, timeout=(8, 45))
        if response.status_code == 404:
            raise RuntimeError(f"Forge尚未部署只读接口：/{path}")
        response.raise_for_status()
        try:
            payload = response.json()
        except ValueError as exc:
            raise RuntimeError(f"Forge接口 /{path} 未返回JSON") from exc
        records, deleted, next_cursor, has_more, sync_cursor = unpack_page(payload)
        all_records.extend(records)
        all_deleted.extend(deleted)
        final_sync_cursor = sync_cursor or final_sync_cursor
        if next_cursor and next_cursor not in seen_cursors:
            seen_cursors.add(next_cursor)
            page_cursor = next_cursor
            page += 1
            continue
        if has_more and records:
            page += 1
            page_cursor = ""
            continue
        break
    else:
        raise RuntimeError("Forge分页超过安全上限，已停止同步")
    return all_records, all_deleted, final_sync_cursor


def ingest_records(records: Iterable[dict[str, Any]], deleted_ids: Iterable[Any] = ()) -> dict[str, Any]:
    conn = get_db()
    received = changed = deleted = 0
    max_updated = ""
    errors: list[str] = []
    try:
        conn.execute("BEGIN IMMEDIATE")
        for raw in records:
            received += 1
            try:
                record = normalize_record(raw)
                was_changed, was_deleted = upsert_record(conn, record)
                changed += int(was_changed)
                deleted += int(was_changed and was_deleted)
                max_updated = max(max_updated, record["source_updated_at"])
            except (ValueError, TypeError) as exc:
                errors.append(str(exc))
        for item in deleted_ids:
            if isinstance(item, dict):
                record_id = pick(item, "complaint_record_id", "record_id", "id")
                updated_at = pick(item, "source_updated_at", "updated_at")
            else:
                record_id, updated_at = item, ""
            was_changed, _ = mark_deleted(conn, clean(record_id), clean(updated_at))
            changed += int(was_changed)
            deleted += int(was_changed)
            max_updated = max(max_updated, normalize_datetime(updated_at))
        if changed:
            conn.execute(
                "UPDATE complaint_audit_sync_cursor SET data_version=data_version+1 WHERE source_key=?", (SYNC_KEY,)
            )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {
        "received": received, "changed": changed, "deleted": deleted,
        "max_updated": max_updated, "errors": errors,
    }


def excel_row_to_record(row: dict[str, Any]) -> dict[str, Any]:
    parts = split_reason_path(row.get("原因分类别名"))
    result = parts[0] if parts and parts[0] in {"客诉成立", "客诉不成立"} else ""
    branch = parts[1] if len(parts) > 1 else ""
    normalized = dict(row)
    normalized.update(
        {
            "complaint_record_id": identifier(row.get("ID")),
            "inspect_id": identifier(row.get("稽查单ID")),
            "ticket_id": identifier(row.get("工单ID")),
            "shop_id": identifier(row.get("店铺ID")),
            "shop_name": "",
            "brand": clean(row.get("KA品牌名称")),
            "inspect_status": clean(row.get("稽查状态")),
            "inspect_completed_at": row.get("完成处理时间"),
            "complaint_result": "是" if result == "客诉成立" else "否" if result == "客诉不成立" else "",
            "review_status": row.get("审核状态"),
            "complained_agent": clean(row.get("昵称")),
            "reason_node_id": "",
            "reason_path": normalized_reason_path(row.get("原因分类别名")),
            "responsibility_subject": EXCEL_RESPONSIBILITY_BRANCH_MAP.get(branch, ""),
            "reason_description": strip_html(row.get("客诉原因描述")),
            "inspect_result": "",
            "abnormal_state": "",
            "abnormal_type": "",
            "handled_at": row.get("完成处理时间"),
            "reviewed_at": row.get("完成审核时间"),
            "source_updated_at": row.get("最后修改日期"),
            "deleted": False,
            "feedback_source": clean(row.get("反馈来源")),
            "feedback_at": normalize_datetime(row.get("反馈时间")),
            "event_at": normalize_datetime(row.get("事件发生时间")),
            "complaint_category": normalized_reason_path(row.get("客诉分类别名")),
            "repeat_report": clean(row.get("商家是否重复发起")),
            "duty_related": clean(row.get("是否值守客服相关")),
            "work_order_status": clean(row.get("工单状态")),
            "processing_status": clean(row.get("处理状态")),
            "time_basis": EXCEL_TIME_BASIS,
        }
    )
    return normalized


def import_excel_workbook(source: Any, filename: str = "Forge导出.xlsx") -> dict[str, Any]:
    started = datetime.now()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        row_iter = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(row_iter, ())]
        missing = sorted(EXCEL_REQUIRED_COLUMNS - set(headers))
        if missing:
            raise ValueError("导出文件缺少必要字段：" + "、".join(missing))
        records: list[dict[str, Any]] = []
        seen_ids: set[str] = set()
        duplicate_source_rows = 0
        missing_completed_inspected = 0
        for values in row_iter:
            if not any(value is not None for value in values):
                continue
            row = dict(zip(headers, values))
            source_id = identifier(row.get("ID"))
            if source_id in seen_ids:
                duplicate_source_rows += 1
            elif source_id:
                seen_ids.add(source_id)
            if clean(row.get("稽查状态")) == "已稽查" and not normalize_datetime(row.get("完成处理时间")):
                missing_completed_inspected += 1
            records.append(excel_row_to_record(row))
    finally:
        workbook.close()

    batch = ingest_records(records)
    finished = datetime.now()
    duration_ms = int((finished - started).total_seconds() * 1000)
    skipped = len(batch["errors"])
    message = (
        f"Excel导入{batch['received']}条，更新{batch['changed']}条，"
        f"跳过{skipped}条；按{EXCEL_TIME_BASIS}统计"
    )
    conn = get_db()
    conn.execute(
        "UPDATE complaint_audit_sync_cursor SET last_source_updated_at=?,last_success_at=? WHERE source_key=?",
        (batch["max_updated"], finished.isoformat(timespec="seconds"), SYNC_KEY),
    )
    conn.execute(
        "INSERT INTO complaint_audit_sync_runs "
        "(mode,started_at,finished_at,status,received_count,changed_count,deleted_count,duration_ms,message) "
        "VALUES (?,?,?,?,?,?,?,?,?)",
        (
            "excel_import", started.isoformat(timespec="seconds"), finished.isoformat(timespec="seconds"),
            "ok", batch["received"], batch["changed"], 0, duration_ms, message,
        ),
    )
    conn.commit()
    conn.close()
    AUTO_STATE.update(
        {
            "running": False, "message": message, "last_finished_at": finished.isoformat(timespec="seconds"),
            "last_error": None, "next_run_at": None,
        }
    )
    return {
        **batch,
        "filename": Path(filename).name,
        "duration_ms": duration_ms,
        "duplicate_source_rows": duplicate_source_rows,
        "missing_completed_inspected": missing_completed_inspected,
        "skipped": skipped,
        "time_basis": EXCEL_TIME_BASIS,
        "message": message,
    }


def sync_once(force_backfill: bool = False) -> dict[str, Any]:
    if not SYNC_LOCK.acquire(blocking=False):
        raise RuntimeError("已有客诉增量同步正在运行")
    started = datetime.now()
    mode = "incremental"
    run_id = None
    AUTO_STATE.update({"running": True, "message": "正在检查客诉稽查增量…", "last_started_at": now_iso(), "last_error": None})
    try:
        init_db()
        conn = get_db()
        cursor_row = conn.execute(
            "SELECT * FROM complaint_audit_sync_cursor WHERE source_key=?", (SYNC_KEY,)
        ).fetchone()
        record_count = conn.execute("SELECT COUNT(*) FROM complaint_audit_records").fetchone()[0]
        backfill_completed = bool(cursor_row and clean(cursor_row["last_success_at"]))
        mode = "backfill" if force_backfill or record_count == 0 or not backfill_completed else "incremental"
        run_id = conn.execute(
            "INSERT INTO complaint_audit_sync_runs (mode,started_at,status,message) VALUES (?,?,?,?)",
            (mode, started.isoformat(timespec="seconds"), "running", "同步中"),
        ).lastrowid
        conn.commit()
        conn.close()

        totals = {"received": 0, "changed": 0, "deleted": 0, "max_updated": "", "errors": []}
        new_cursor = clean(cursor_row["cursor"]) if cursor_row else ""
        if mode == "backfill":
            AUTO_STATE["message"] = "正在回溯近6个月客诉稽查记录…"
            records, deleted_ids, returned_cursor = fetch_pages(
                BACKFILL_PATH,
                {"completed_from": six_month_start().isoformat(), "completed_to": date.today().isoformat()},
            )
            batch = ingest_records(records, deleted_ids)
            new_cursor = returned_cursor or new_cursor
            for key in totals:
                totals[key] = max(totals[key], batch[key]) if key == "max_updated" else totals[key] + batch[key] if key != "errors" else totals[key] + batch[key]
        else:
            since = subtract_minutes(clean(cursor_row["last_source_updated_at"]), OVERLAP_MINUTES)
            records, deleted_ids, returned_cursor = fetch_pages(
                CHANGES_PATH, {"cursor": new_cursor, "modified_since": since}
            )
            batch = ingest_records(records, deleted_ids)
            new_cursor = returned_cursor or new_cursor
            totals = batch

        today_text = date.today().isoformat()
        last_reconcile = clean(cursor_row["last_reconcile_date"]) if cursor_row else ""
        if mode == "backfill" or last_reconcile != today_text:
            AUTO_STATE["message"] = "正在核对本月和上月客诉稽查记录…"
            start = shift_month(date.today().replace(day=1), -1)
            records, deleted_ids, returned_cursor = fetch_pages(
                BACKFILL_PATH, {"completed_from": start.isoformat(), "completed_to": date.today().isoformat(), "reconcile": 1}
            )
            batch = ingest_records(records, deleted_ids)
            new_cursor = returned_cursor or new_cursor
            totals["received"] += batch["received"]
            totals["changed"] += batch["changed"]
            totals["deleted"] += batch["deleted"]
            totals["max_updated"] = max(totals["max_updated"], batch["max_updated"])
            totals["errors"].extend(batch["errors"])
            last_reconcile = today_text

        finished = datetime.now()
        duration_ms = int((finished - started).total_seconds() * 1000)
        conn = get_db()
        existing_watermark = clean(cursor_row["last_source_updated_at"]) if cursor_row else ""
        watermark = max(existing_watermark, totals["max_updated"])
        conn.execute(
            "UPDATE complaint_audit_sync_cursor SET cursor=?,last_source_updated_at=?,last_success_at=?,last_reconcile_date=? WHERE source_key=?",
            (new_cursor, watermark, finished.isoformat(timespec="seconds"), last_reconcile, SYNC_KEY),
        )
        message = f"收到{totals['received']}条，更新{totals['changed']}条，删除{totals['deleted']}条"
        if totals["errors"]:
            message += f"；跳过{len(totals['errors'])}条字段异常记录"
        conn.execute(
            "UPDATE complaint_audit_sync_runs SET finished_at=?,status='ok',received_count=?,changed_count=?,deleted_count=?,duration_ms=?,message=? WHERE id=?",
            (finished.isoformat(timespec="seconds"), totals["received"], totals["changed"], totals["deleted"], duration_ms, message, run_id),
        )
        conn.commit()
        conn.close()
        AUTO_STATE.update({"running": False, "message": message, "last_finished_at": finished.isoformat(timespec="seconds"), "last_error": None})
        return {**totals, "mode": mode, "duration_ms": duration_ms, "message": message}
    except Exception as exc:
        finished = datetime.now()
        duration_ms = int((finished - started).total_seconds() * 1000)
        if run_id is not None:
            conn = get_db()
            conn.execute(
                "UPDATE complaint_audit_sync_runs SET finished_at=?,status='error',duration_ms=?,message=? WHERE id=?",
                (finished.isoformat(timespec="seconds"), duration_ms, str(exc), run_id),
            )
            conn.commit()
            conn.close()
        AUTO_STATE.update({"running": False, "message": "同步失败", "last_finished_at": finished.isoformat(timespec="seconds"), "last_error": str(exc)})
        raise
    finally:
        SYNC_LOCK.release()


def auto_sync_loop() -> None:
    while True:
        try:
            sync_once()
        except Exception:
            pass
        next_run = datetime.now() + timedelta(seconds=SYNC_INTERVAL_SECONDS)
        AUTO_STATE["next_run_at"] = next_run.isoformat(timespec="seconds")
        time.sleep(max(1, (next_run - datetime.now()).total_seconds()))


def start_auto_sync() -> None:
    global AUTO_SYNC_STARTED
    if AUTO_SYNC_STARTED or source_mode() == "excel":
        return
    AUTO_SYNC_STARTED = True
    threading.Thread(target=auto_sync_loop, daemon=True, name="complaint-audit-sync").start()


def safe_pct(part: int | float, total: int | float) -> float:
    return round(part * 100 / total, 1) if total else 0.0


def source_url(source_record_id: str) -> str:
    template = os.environ.get(
        "COMPLAINT_AUDIT_FORGE_DETAIL_URL_TEMPLATE",
        "https://forge.weidian.com/admin/hw7cth9v5sk/popups/8quua2pje4s/filterbytk/{record_id}?authenticator=wechat",
    )
    return template.replace("{record_id}", source_record_id)


def classification_for(rows: list[dict[str, Any]]) -> str:
    approved = [row for row in rows if row["review_status"] == "审核通过"]
    results = {row["complaint_result"] for row in approved if row["complaint_result"] in {"是", "否"}}
    if len(results) > 1:
        return "结果冲突"
    if results == {"是"}:
        return "客诉成立"
    if results == {"否"}:
        return "客诉不成立"
    if approved:
        return "未判定"
    if any(row["review_status"] == "待审核" for row in rows):
        return "待审核"
    if any(row["review_status"] == "审核退回" for row in rows):
        return "审核退回"
    return "未审核"


def aggregate_inspection(rows: list[dict[str, Any]]) -> dict[str, Any]:
    ordered = sorted(rows, key=lambda row: (row.get("source_updated_at", ""), row.get("reviewed_at", "")), reverse=True)
    latest = ordered[0]
    approved = [row for row in ordered if row["review_status"] == "审核通过"]
    primary = approved[0] if approved else latest
    try:
        source_payload = json.loads(primary.get("source_payload") or "{}")
    except (TypeError, json.JSONDecodeError):
        source_payload = {}
    classification = classification_for(ordered)
    subjects = {row["responsibility_subject"] for row in approved if row["responsibility_subject"]}
    responsibility_conflict = len(subjects) > 1
    responsibility = next(iter(subjects)) if len(subjects) == 1 else "责任待确认"
    agents = sorted({row["complained_agent"] for row in ordered if row["complained_agent"]})
    paths = sorted({row["reason_path"] for row in approved or ordered if row["reason_path"]})
    primary_path = primary.get("reason_path", "")
    reason_levels = split_reason_path(primary_path)
    nodes = sorted({row["reason_node_id"] for row in approved or ordered if row["reason_node_id"]})
    conclusions = [row["inspect_result"] for row in ordered if row["inspect_result"]]
    abnormalities = [row["abnormal_state"] for row in ordered if row["abnormal_state"]]
    abnormal_types = sorted({row["abnormal_type"] for row in ordered if row["abnormal_type"]})
    descriptions = [row["reason_description"] for row in approved or ordered if row["reason_description"]]
    record_ids = [row["source_record_id"] for row in ordered]
    return {
        "inspect_id": latest["inspect_id"], "ticket_id": latest["ticket_id"],
        "shop_id": latest["shop_id"], "shop_name": latest["shop_name"], "brand": latest["brand"],
        "inspect_status": latest["inspect_status"], "inspect_completed_at": latest["inspect_completed_at"],
        "classification": classification,
        "complaint_result": "是" if classification == "客诉成立" else "否" if classification == "客诉不成立" else "",
        "review_status": "审核通过" if approved else latest["review_status"],
        "complained_agent": "、".join(agents), "complained_agents": agents,
        "responsibility_subject": responsibility, "responsibility_conflict": responsibility_conflict,
        "reason_node_id": "、".join(nodes), "reason_path": "；".join(paths),
        "primary_reason_path": primary_path, "reason_levels": reason_levels,
        "reason_description": descriptions[0] if descriptions else "",
        "inspect_result": conclusions[0] if conclusions else "未填写稽查结论",
        "abnormal_state": abnormalities[0] if abnormalities else "未填写",
        "abnormal_type": "、".join(abnormal_types) if abnormal_types else "未填写",
        "handled_at": latest["handled_at"], "reviewed_at": latest["reviewed_at"],
        "source_updated_at": max(row["source_updated_at"] for row in ordered),
        "record_count": len(ordered), "approved_record_count": len(approved),
        "multiple_records": len(ordered) > 1, "source_record_ids": record_ids,
        "source_url": source_url(record_ids[0]),
        "feedback_source": clean(source_payload.get("feedback_source") or source_payload.get("反馈来源")),
        "feedback_at": normalize_datetime(source_payload.get("feedback_at") or source_payload.get("反馈时间")),
        "event_at": normalize_datetime(source_payload.get("event_at") or source_payload.get("事件发生时间")),
        "complaint_category": normalized_reason_path(source_payload.get("complaint_category") or source_payload.get("客诉分类别名")),
        "repeat_report": clean(source_payload.get("repeat_report") or source_payload.get("商家是否重复发起")),
        "duty_related": clean(source_payload.get("duty_related") or source_payload.get("是否值守客服相关")),
        "time_basis": clean(source_payload.get("time_basis")) or "稽查完成时间",
    }


def load_aggregates(month: str) -> list[dict[str, Any]]:
    start = month_start(month)
    end = month_end(start)
    conn = get_db()
    rows = [dict(row) for row in conn.execute(
        "SELECT * FROM complaint_audit_records "
        "WHERE deleted=0 AND inspect_status='已稽查' AND substr(inspect_completed_at,1,10) BETWEEN ? AND ? "
        "ORDER BY inspect_completed_at,source_updated_at",
        (start.isoformat(), end.isoformat()),
    )]
    conn.close()
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["inspect_id"]].append(row)
    return [aggregate_inspection(group) for group in grouped.values()]


def month_coverage(month: str) -> dict[str, Any]:
    start = month_start(month)
    end = month_end(start)
    conn = get_db()
    row = conn.execute(
        "SELECT MIN(substr(inspect_completed_at,1,10)),MAX(substr(inspect_completed_at,1,10)),COUNT(*) "
        "FROM complaint_audit_records WHERE deleted=0 AND inspect_status='已稽查' "
        "AND substr(inspect_completed_at,1,10) BETWEEN ? AND ?",
        (start.isoformat(), end.isoformat()),
    ).fetchone()
    missing_completed = conn.execute(
        "SELECT COUNT(*) FROM complaint_audit_records WHERE deleted=0 AND inspect_status='已稽查' "
        "AND COALESCE(inspect_completed_at,'')=''"
    ).fetchone()[0]
    conn.close()
    first_day, last_day, row_count = row
    expected_end = min(end, date.today())
    complete = bool(first_day == start.isoformat() and last_day and last_day >= expected_end.isoformat())
    return {
        "month": month, "first_day": first_day, "last_day": last_day, "row_count": row_count,
        "complete_to_date": complete, "missing_completed_records": missing_completed,
    }


def query_filters() -> dict[str, Any]:
    month = clean(request.args.get("month")) or date.today().strftime("%Y-%m")
    selected = month_start(month)
    if selected < six_month_start() or selected > date.today().replace(day=1):
        raise ValueError("仅支持查看近6个月数据")
    return {
        "month": month,
        "date_from": clean(request.args.get("date_from")), "date_to": clean(request.args.get("date_to")),
        "review_status": clean(request.args.get("review_status")),
        "complaint_result": clean(request.args.get("complaint_result")),
        "responsibility_subject": clean(request.args.get("responsibility_subject")),
        "agent": clean(request.args.get("agent")), "brand": clean(request.args.get("brand")),
        "root_cause": clean(request.args.get("root_cause")),
        "complaint_category": clean(request.args.get("complaint_category")),
        "description_theme": clean(request.args.get("description_theme")),
        "description_group": clean(request.args.get("description_group")),
        "reason_description": normalized_description(request.args.get("reason_description")),
        "q": clean(request.args.get("q")).lower(), "bucket": clean(request.args.get("bucket")),
    }


def matches_filters(record: dict[str, Any], filters: dict[str, Any]) -> bool:
    completed_day = record["inspect_completed_at"][:10]
    if filters["date_from"] and completed_day < filters["date_from"]:
        return False
    if filters["date_to"] and completed_day > filters["date_to"]:
        return False
    review = filters["review_status"]
    if review and not (review == record["review_status"] or review == record["classification"]):
        return False
    result = filters["complaint_result"]
    if result and result != record["classification"] and result != record["complaint_result"]:
        return False
    if filters["responsibility_subject"] and filters["responsibility_subject"] != record["responsibility_subject"]:
        return False
    if filters["agent"] and filters["agent"] not in record["complained_agents"]:
        return False
    if filters["brand"] and filters["brand"] != record["brand"]:
        return False
    if filters["root_cause"]:
        cause_path = " / ".join(record.get("reason_levels", [])[1:])
        if filters["root_cause"] not in record.get("reason_levels", []) and filters["root_cause"] not in {
            cause_path, record.get("primary_reason_path", ""),
        } and filters["root_cause"] not in cause_path:
            return False
    if filters["complaint_category"] and filters["complaint_category"] != record.get("complaint_category", ""):
        return False
    description_theme, description_group = description_theme_for(record.get("reason_description"))
    if filters["description_theme"] and filters["description_theme"] != description_theme:
        return False
    if filters["description_group"] and filters["description_group"] != description_group:
        return False
    record_description = normalized_description(record.get("reason_description")) or "未填写"
    if filters["reason_description"] and filters["reason_description"] != record_description:
        return False
    bucket_map = {
        "inspected": None, "approved": {"客诉成立", "客诉不成立", "未判定", "结果冲突"},
        "upheld": {"客诉成立"}, "not_upheld": {"客诉不成立"}, "undetermined": {"未判定"},
        "pending": {"待审核"}, "returned": {"审核退回"}, "conflict": {"结果冲突"},
    }
    bucket = filters["bucket"]
    if bucket and bucket in bucket_map and bucket_map[bucket] is not None and record["classification"] not in bucket_map[bucket]:
        return False
    if filters["q"]:
        haystack = " ".join(clean(record.get(key)) for key in (
            "inspect_id", "ticket_id", "shop_id", "shop_name", "brand", "complained_agent",
            "reason_node_id", "reason_path", "reason_description", "inspect_result", "abnormal_type",
        )).lower()
        if filters["q"] not in haystack:
            return False
    return True


def filter_records(records: list[dict[str, Any]], filters: dict[str, Any]) -> list[dict[str, Any]]:
    return [record for record in records if matches_filters(record, filters)]


def summary_counts(records: list[dict[str, Any]]) -> dict[str, Any]:
    counts = Counter(record["classification"] for record in records)
    upheld = counts["客诉成立"]
    not_upheld = counts["客诉不成立"]
    approved = upheld + not_upheld + counts["未判定"] + counts["结果冲突"]
    return {
        "inspected": len(records), "approved": approved, "upheld": upheld, "not_upheld": not_upheld,
        "upheld_rate": safe_pct(upheld, upheld + not_upheld), "undetermined": counts["未判定"],
        "pending": counts["待审核"], "returned": counts["审核退回"], "conflict": counts["结果冲突"],
        "unreviewed": counts["未审核"],
        "agent_count": len({agent for record in records for agent in record["complained_agents"]}),
    }


def trend_rows(records: list[dict[str, Any]], month: str, granularity: str) -> list[dict[str, Any]]:
    start = month_start(month)
    end = month_end(start)
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    if granularity == "week":
        for record in records:
            day = date.fromisoformat(record["inspect_completed_at"][:10])
            monday = day - timedelta(days=day.weekday())
            grouped[monday.isoformat()].append(record)
        keys = sorted(grouped)
        rows = []
        for key in keys:
            monday = date.fromisoformat(key)
            week_end = min(monday + timedelta(days=6), end)
            counts = summary_counts(grouped[key])
            rows.append({"key": key, "label": f"{max(monday,start).strftime('%m/%d')}-{week_end.strftime('%m/%d')}", **counts})
        return rows
    for record in records:
        grouped[record["inspect_completed_at"][:10]].append(record)
    day = start
    while day <= end:
        grouped.setdefault(day.isoformat(), [])
        day += timedelta(days=1)
    return [
        {"key": key, "label": date.fromisoformat(key).strftime("%m/%d"), **summary_counts(grouped[key])}
        for key in sorted(grouped)
    ]


def responsibility_analysis(records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    upheld = [record for record in records if record["classification"] == "客诉成立"]
    dist = Counter(record["responsibility_subject"] for record in upheld)
    distribution = [
        {"subject": subject, "count": dist[subject], "ratio": safe_pct(dist[subject], len(upheld)),
         "filters": {"bucket": "upheld", "responsibility_subject": subject}}
        for subject in RESPONSIBILITY_SUBJECTS
    ]
    by_agent: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        for agent in record["complained_agents"]:
            by_agent[agent].append(record)
    rankings = []
    for agent, agent_records in by_agent.items():
        decided = [row for row in agent_records if row["classification"] in {"客诉成立", "客诉不成立"}]
        complaint_upheld = sum(row["classification"] == "客诉成立" for row in decided)
        formal_responsibility = sum(
            row["classification"] == "客诉成立" and row["responsibility_subject"] == "人工客服责任"
            for row in agent_records
        )
        rankings.append({
            "agent": agent, "complained_count": len(agent_records),
            "formal_responsibility_count": formal_responsibility,
            "complaint_upheld_count": complaint_upheld, "decided_sample": len(decided),
            "upheld_rate": safe_pct(complaint_upheld, len(decided)),
        })
    rankings.sort(key=lambda row: (-row["formal_responsibility_count"], -row["complaint_upheld_count"], -row["complained_count"], row["agent"]))
    return distribution, rankings


def ranked_rows(counter: Counter, total: int, key_name: str, limit: int = 15) -> list[dict[str, Any]]:
    return [
        {key_name: key, "count": count, "ratio": safe_pct(count, total)}
        for key, count in counter.most_common(limit)
    ]


def root_cause_analysis(records: list[dict[str, Any]]) -> dict[str, Any]:
    upheld = [record for record in records if record["classification"] == "客诉成立"]
    path_counter: Counter = Counter()
    branch_counter: Counter = Counter()
    mechanism_counter: Counter = Counter()
    detail_counter: Counter = Counter()
    category_counter: Counter = Counter()
    manual_mechanism: Counter = Counter()
    manual_detail: Counter = Counter()
    for record in upheld:
        levels = record.get("reason_levels", [])
        branch = levels[1] if len(levels) > 1 else "责任路径未填写"
        mechanism = levels[2] if len(levels) > 2 else "三级原因未填写"
        detail = " / ".join(levels[3:]) if len(levels) > 3 else "细分原因未填写"
        cause_path = " / ".join(levels[1:]) if len(levels) > 1 else "责任路径未填写"
        category = record.get("complaint_category") or "客诉分类未填写"
        path_counter[cause_path] += 1
        branch_counter[branch] += 1
        mechanism_counter[mechanism] += 1
        detail_counter[detail] += 1
        category_counter[category] += 1
        if record["responsibility_subject"] == "人工客服责任":
            manual_mechanism[mechanism] += 1
            manual_detail[detail] += 1
    path_rows = ranked_rows(path_counter, len(upheld), "path", 20)
    for row in path_rows:
        row["filters"] = {"bucket": "upheld", "root_cause": row["path"]}
    branch_rows = ranked_rows(branch_counter, len(upheld), "name", 12)
    for row in branch_rows:
        row["filters"] = {"bucket": "upheld", "root_cause": row["name"]}
    mechanism_rows = ranked_rows(mechanism_counter, len(upheld), "name", 15)
    for row in mechanism_rows:
        row["filters"] = {"bucket": "upheld", "root_cause": row["name"]}
    detail_rows = ranked_rows(detail_counter, len(upheld), "name", 20)
    for row in detail_rows:
        row["filters"] = {"bucket": "upheld", "root_cause": row["name"]}
    manual_total = sum(manual_mechanism.values())
    manual_mechanism_rows = ranked_rows(manual_mechanism, manual_total, "name", 12)
    manual_detail_rows = ranked_rows(manual_detail, manual_total, "name", 20)
    for row in manual_mechanism_rows + manual_detail_rows:
        row["filters"] = {
            "bucket": "upheld", "responsibility_subject": "人工客服责任", "root_cause": row["name"],
        }
    category_rows = ranked_rows(category_counter, len(upheld), "name", 15)
    for row in category_rows:
        row["filters"] = {"bucket": "upheld", "complaint_category": row["name"]}
    return {
        "upheld_total": len(upheld),
        "paths": path_rows,
        "branches": branch_rows,
        "mechanisms": mechanism_rows,
        "details": detail_rows,
        "manual_total": manual_total,
        "manual_mechanisms": manual_mechanism_rows,
        "manual_details": manual_detail_rows,
        "categories": category_rows,
        "quality": {
            "multiple_record_inspections": sum(record["multiple_records"] for record in records),
            "extra_source_records": sum(max(0, record["record_count"] - 1) for record in records),
            "result_conflicts": sum(record["classification"] == "结果冲突" for record in records),
            "missing_conclusion": sum(record["inspect_result"] == "未填写稽查结论" for record in records),
            "responsibility_pending": sum(
                record["classification"] == "客诉成立" and record["responsibility_subject"] == "责任待确认"
                for record in records
            ),
        },
    }


def not_upheld_cause_analysis(records: list[dict[str, Any]]) -> dict[str, Any]:
    """按正式分类路径拆解审核通过且客诉不成立的唯一稽查单。"""
    not_upheld = [record for record in records if record["classification"] == "客诉不成立"]
    branch_counter: Counter = Counter()
    mechanism_counter: Counter = Counter()
    detail_counter: Counter = Counter()
    path_counter: Counter = Counter()
    category_counter: Counter = Counter()
    details_by_branch: dict[str, Counter] = defaultdict(Counter)
    for record in not_upheld:
        levels = record.get("reason_levels", [])
        branch = levels[1] if len(levels) > 1 else "不成立原因未填写"
        mechanism = levels[2] if len(levels) > 2 else "三级原因未填写"
        detail = " / ".join(levels[3:]) if len(levels) > 3 else "细分原因未填写"
        cause_path = " / ".join(levels[1:]) if len(levels) > 1 else "不成立原因未填写"
        category = record.get("complaint_category") or "客诉分类未填写"
        branch_counter[branch] += 1
        mechanism_counter[mechanism] += 1
        detail_counter[detail] += 1
        details_by_branch[branch][detail] += 1
        path_counter[cause_path] += 1
        category_counter[category] += 1

    def cause_rows(counter: Counter, key_name: str = "name", limit: int = 15) -> list[dict[str, Any]]:
        rows = ranked_rows(counter, len(not_upheld), key_name, limit)
        for row in rows:
            row["filters"] = {"bucket": "not_upheld", "root_cause": row[key_name]}
        return rows

    branches = cause_rows(branch_counter, limit=12)
    mechanisms = cause_rows(mechanism_counter, limit=15)
    details = cause_rows(detail_counter, limit=20)
    paths = cause_rows(path_counter, key_name="path", limit=20)
    categories = ranked_rows(category_counter, len(not_upheld), "name", 15)
    for row in categories:
        row["filters"] = {"bucket": "not_upheld", "complaint_category": row["name"]}

    def count_for(counter: Counter, name: str) -> int:
        return int(counter.get(name, 0))

    no_abnormal = count_for(branch_counter, "核实结果无异常")
    repeated = count_for(branch_counter, "重复反馈")
    manager = count_for(branch_counter, "服务管家原因")
    multi_channel = count_for(detail_counter, "多渠道重复提交")
    verified_repeat = count_for(detail_counter, "稽查已核实商家重复反馈")
    analyst_facts: list[dict[str, Any]] = []
    top_two_total = no_abnormal + repeated
    if top_two_total:
        analyst_facts.append({
            "text": f"“核实结果无异常”和“重复反馈”合计{top_two_total}张，占全部客诉不成立单{safe_pct(top_two_total, len(not_upheld))}%；其余正式原因共{len(not_upheld) - top_two_total}张。",
            "filters": {"bucket": "not_upheld"},
        })
    if no_abnormal:
        no_abnormal_details = ranked_rows(details_by_branch["核实结果无异常"], no_abnormal, "name", 3)
        detail_text = "、".join(f"“{row['name']}”{row['count']}张" for row in no_abnormal_details[:3])
        analyst_facts.append({
            "text": f"{no_abnormal}张（{safe_pct(no_abnormal, len(not_upheld))}%）正式分类为“核实结果无异常”"
                    + (f"，主要细分为{detail_text}。" if detail_text else "。"),
            "filters": {"bucket": "not_upheld", "root_cause": "核实结果无异常"},
        })
    if repeated:
        parts = []
        if verified_repeat:
            parts.append(f"“稽查已核实商家重复反馈”{verified_repeat}张")
        if multi_channel:
            parts.append(f"“多渠道重复提交”{multi_channel}张")
        analyst_facts.append({
            "text": f"{repeated}张（{safe_pct(repeated, len(not_upheld))}%）正式分类为“重复反馈”"
                    + (f"，其中{'、'.join(parts)}。" if parts else "。"),
            "filters": {"bucket": "not_upheld", "root_cause": "重复反馈"},
        })
    if manager:
        manager_details = ranked_rows(details_by_branch["服务管家原因"], manager, "name", 5)
        detail_text = "、".join(f"“{row['name']}”{row['count']}张" for row in manager_details)
        analyst_facts.append({
            "text": f"{manager}张（{safe_pct(manager, len(not_upheld))}%）正式分类为“服务管家原因”"
                    + (f"，细分为{detail_text}。" if detail_text else "。"),
            "filters": {"bucket": "not_upheld", "root_cause": "服务管家原因"},
        })
    if categories:
        top = categories[0]
        analyst_facts.append({
            "text": f"客诉场景中数量最多的是“{top['name']}”，共{top['count']}张，占不成立单{top['ratio']}%；该比例仅反映正式客诉分类。",
            "filters": top["filters"],
        })
    return {
        "not_upheld_total": len(not_upheld),
        "branches": branches,
        "mechanisms": mechanisms,
        "details": details,
        "paths": paths,
        "categories": categories,
        "analyst_facts": analyst_facts,
        "repeat_summary": {
            "repeat_count": repeated,
            "repeat_ratio": safe_pct(repeated, len(not_upheld)),
            "verified_repeat_count": verified_repeat,
            "multi_channel_count": multi_channel,
            "multi_channel_ratio": safe_pct(multi_channel, len(not_upheld)),
        },
    }


def not_upheld_description_analysis(records: list[dict[str, Any]]) -> dict[str, Any]:
    """读取原因描述原文，按可审计规则做主题聚合，并保留原文高频项。"""
    not_upheld = [record for record in records if record["classification"] == "客诉不成立"]
    theme_counter: Counter = Counter()
    group_counter: Counter = Counter()
    original_counter: Counter = Counter()
    examples: dict[str, list[dict[str, str]]] = defaultdict(list)
    for record in not_upheld:
        text = normalized_description(record.get("reason_description")) or "未填写"
        theme, group = description_theme_for(text)
        theme_counter[theme] += 1
        group_counter[group] += 1
        original_counter[text] += 1
        if len(examples[theme]) < 3 and all(item["text"] != text for item in examples[theme]):
            examples[theme].append({"text": text, "inspect_id": record["inspect_id"]})

    theme_rows = ranked_rows(theme_counter, len(not_upheld), "name", 20)
    for row in theme_rows:
        row["group"] = next(
            group for theme, group, _pattern in DESCRIPTION_THEME_RULES if theme == row["name"]
        ) if row["name"] != "其他未归类原因描述" else "其他未归类"
        row["examples"] = examples[row["name"]]
        row["filters"] = {"bucket": "not_upheld", "description_theme": row["name"]}
    group_rows = ranked_rows(group_counter, len(not_upheld), "name", 10)
    for row in group_rows:
        row["filters"] = {"bucket": "not_upheld", "description_group": row["name"]}
    original_rows = ranked_rows(original_counter, len(not_upheld), "name", 30)
    for row in original_rows:
        row["filters"] = {"bucket": "not_upheld", "reason_description": row["name"]}

    evidence_count = group_counter["包含明确核实结果或状态"]
    request_only_count = group_counter["仅描述诉求或处理动作"]
    vague_count = group_counter["描述缺失或过于笼统"]
    unresolved_count = group_counter["描述仍含异常或待处理"]
    other_count = group_counter["其他未归类"]
    facts = [
        {
            "text": f"{evidence_count}张（{safe_pct(evidence_count, len(not_upheld))}%）原因描述包含明确核实结果或状态；该数字来自固定文本规则归类，点击可逐单核对原文。",
            "filters": {"bucket": "not_upheld", "description_group": "包含明确核实结果或状态"},
        },
        {
            "text": f"{request_only_count}张（{safe_pct(request_only_count, len(not_upheld))}%）仅写核实、视频、赔付、介入或通知动作，未在原因描述中写明最终不成立依据。",
            "filters": {"bucket": "not_upheld", "description_group": "仅描述诉求或处理动作"},
        },
    ]
    if vague_count or other_count:
        facts.append({
            "text": f"另有描述缺失或过于笼统{vague_count}张、其他未归类描述{other_count}张，当前不强行判断具体原因。",
            "filters": {"bucket": "not_upheld", "description_group": "描述缺失或过于笼统"},
        })
    if unresolved_count:
        facts.append({
            "text": f"{unresolved_count}张原因描述仍含漏付、少付、偷盗、未支付等异常或待处理表述，建议结合稽查结论和原单复核，不宜只凭“不成立”标签下结论。",
            "filters": {"bucket": "not_upheld", "description_group": "描述仍含异常或待处理"},
        })
    return {
        "not_upheld_total": len(not_upheld),
        "groups": group_rows,
        "themes": theme_rows,
        "originals": original_rows,
        "facts": facts,
        "method": "原因描述按固定可审计规则归类；原文始终保留，无法明确归类时进入未归类或待复核。",
    }


def fact_summaries(
    records: list[dict[str, Any]], previous: list[dict[str, Any]],
    coverage: dict[str, Any], previous_coverage: dict[str, Any],
) -> list[dict[str, Any]]:
    current = summary_counts(records)
    prior = summary_counts(previous)
    facts = [
        {"text": f"本期共有{current['inspected']}张状态为“已稽查”的稽查单，其中{current['approved']}张存在审核通过记录。", "filters": {"bucket": "inspected"}},
        {"text": f"审核通过且结果明确的稽查单共{current['upheld'] + current['not_upheld']}张：客诉成立{current['upheld']}张、客诉不成立{current['not_upheld']}张，成立率{current['upheld_rate']}%。", "filters": {"bucket": "approved"}},
    ]
    exceptions = current["undetermined"] + current["pending"] + current["returned"] + current["conflict"]
    if exceptions:
        facts.append({
            "text": f"当前另有未判定{current['undetermined']}张、待审核{current['pending']}张、审核退回{current['returned']}张、结果冲突待核验{current['conflict']}张，均未混入正式成立或不成立数字。",
            "filters": {},
        })
    if coverage["complete_to_date"] and previous_coverage["complete_to_date"]:
        delta = current["inspected"] - prior["inspected"]
        facts.append({
            "text": f"与上月可比口径相比，本期已稽查单数{'增加' if delta > 0 else '减少' if delta < 0 else '持平'}{abs(delta)}张（本期{current['inspected']}张，上月{prior['inspected']}张）。",
            "filters": {"bucket": "inspected"},
        })
    elif previous_coverage["first_day"]:
        facts.append({
            "text": f"上月数据仅覆盖{previous_coverage['first_day']}至{previous_coverage['last_day']}，当前不生成月环比结论。",
            "filters": {},
        })
    dist, _ = responsibility_analysis(records)
    confirmed = [item for item in dist if item["subject"] != "责任待确认" and item["count"]]
    if confirmed:
        top = max(confirmed, key=lambda item: item["count"])
        facts.append({
            "text": f"审核通过且客诉成立的记录中，“{top['subject']}”已填写{top['count']}张，占全部成立记录的{top['ratio']}%；该比例仅反映正式责任字段。",
            "filters": top["filters"],
        })
    roots = root_cause_analysis(records)
    if roots["mechanisms"]:
        top = roots["mechanisms"][0]
        facts.append({
            "text": f"客诉成立单的正式三级原因中，“{top['name']}”覆盖{top['count']}张，占成立单{top['ratio']}%；该表述仅复述正式分类，不推断因果。",
            "filters": top["filters"],
        })
    if roots["manual_mechanisms"]:
        top = roots["manual_mechanisms"][0]
        facts.append({
            "text": f"明确归入人工客服责任的成立单共{roots['manual_total']}张，其中正式三级原因“{top['name']}”为{top['count']}张，占人工客服责任成立单{top['ratio']}%。",
            "filters": top["filters"],
        })
    not_upheld_causes = not_upheld_cause_analysis(records)
    if not_upheld_causes["branches"]:
        top = not_upheld_causes["branches"][0]
        facts.append({
            "text": f"客诉不成立单的正式原因类型中，“{top['name']}”为{top['count']}张，占不成立单{top['ratio']}%；具体构成见下方不成立原因拆解。",
            "filters": top["filters"],
        })
    description_analysis = not_upheld_description_analysis(records)
    if description_analysis["groups"]:
        top = description_analysis["groups"][0]
        facts.append({
            "text": f"客诉不成立单的原因描述中，“{top['name']}”为{top['count']}张，占不成立单{top['ratio']}%；原文主题与样本见下方深层原因分析。",
            "filters": top["filters"],
        })
    return facts


def sync_status() -> dict[str, Any]:
    init_db()
    mode = source_mode()
    conn = get_db()
    cursor = conn.execute("SELECT * FROM complaint_audit_sync_cursor WHERE source_key=?", (SYNC_KEY,)).fetchone()
    if mode == "excel":
        run = conn.execute(
            "SELECT * FROM complaint_audit_sync_runs WHERE mode='excel_import' ORDER BY id DESC LIMIT 1"
        ).fetchone()
    else:
        run = conn.execute("SELECT * FROM complaint_audit_sync_runs ORDER BY id DESC LIMIT 1").fetchone()
    record_count = conn.execute("SELECT COUNT(*) FROM complaint_audit_records WHERE deleted=0").fetchone()[0]
    source_time = conn.execute("SELECT MAX(source_updated_at) FROM complaint_audit_records WHERE deleted=0").fetchone()[0]
    conn.close()
    last_success = parse_datetime(cursor["last_success_at"]) if cursor else None
    age = max(0, int((datetime.now() - last_success.replace(tzinfo=None)).total_seconds())) if last_success else None
    freshness = "manual" if mode == "excel" and last_success else (
        "normal" if age is not None and age <= 90 else "delayed" if age is not None and age <= 180 else "stale"
    )
    return {
        **AUTO_STATE, "source_updated_at": source_time or None,
        "local_synced_at": cursor["last_success_at"] or None if cursor else None,
        "data_version": cursor["data_version"] if cursor else 0,
        "record_count": record_count, "age_seconds": age, "freshness": freshness,
        "last_run": dict(run) if run else None, "interval_seconds": SYNC_INTERVAL_SECONDS,
        "source_mode": mode, "time_basis": EXCEL_TIME_BASIS if mode == "excel" else "稽查完成时间",
        "upstream_ready": mode == "excel" or not (AUTO_STATE.get("last_error") or "").startswith("Forge尚未部署只读接口"),
    }


def client_addr() -> str:
    return request.remote_addr or "unknown"


def is_loopback() -> bool:
    return client_addr() in {"127.0.0.1", "::1"}


def auth_configured() -> bool:
    return bool(os.environ.get("COMPLAINT_AUDIT_PASSWORD_HASH") and os.environ.get("COMPLAINT_AUDIT_SESSION_SECRET"))


def serializer() -> URLSafeTimedSerializer | None:
    secret = os.environ.get("COMPLAINT_AUDIT_SESSION_SECRET", "")
    return URLSafeTimedSerializer(secret, salt="complaint-audit-session") if secret else None


def current_session() -> dict[str, Any] | None:
    signer = serializer()
    token = request.cookies.get(SESSION_COOKIE, "")
    if not signer or not token:
        return None
    try:
        payload = signer.loads(token, max_age=SESSION_MAX_AGE_SECONDS)
    except (BadSignature, SignatureExpired):
        return None
    return payload if isinstance(payload, dict) and payload.get("authenticated") is True else None


def append_access_log(result: str) -> None:
    init_db()
    conn = get_db()
    conn.execute(
        "INSERT INTO complaint_audit_access_log (attempted_at,client_addr,result) VALUES (?,?,?)",
        (now_iso(), client_addr(), result),
    )
    conn.commit()
    conn.close()


def login_failure_count() -> int:
    cutoff = (datetime.now() - timedelta(minutes=LOGIN_WINDOW_MINUTES)).isoformat(timespec="seconds")
    conn = get_db()
    count = conn.execute(
        "SELECT COUNT(*) FROM complaint_audit_access_log WHERE client_addr=? AND result='失败' AND attempted_at>=?",
        (client_addr(), cutoff),
    ).fetchone()[0]
    conn.close()
    return count


def same_origin() -> bool:
    origin = request.headers.get("Origin", "")
    return not origin or origin in {f"http://{request.host}", f"https://{request.host}"}


def require_auth_response() -> Response:
    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "error": "请先登录客诉稽查看板"}), 401
    return redirect(url_for("complaint_audit.login", next=request.full_path))


def set_session_cookie(response: Response) -> None:
    signer = serializer()
    if not signer:
        return
    token = signer.dumps({"authenticated": True, "csrf": secrets.token_urlsafe(24), "issued_at": now_iso()})
    secure = request.is_secure or os.environ.get("COMPLAINT_AUDIT_COOKIE_SECURE", "").lower() in {"1", "true", "yes"}
    response.set_cookie(
        SESSION_COOKIE, token, max_age=SESSION_MAX_AGE_SECONDS, httponly=True,
        secure=secure, samesite="Lax", path="/",
    )


def validate_csrf() -> bool:
    session = current_session() or {}
    supplied = request.headers.get("X-CSRF-Token") or request.form.get("csrf_token", "")
    return bool(supplied and secrets.compare_digest(clean(supplied), clean(session.get("csrf"))))


def register_app_guards(app: Any) -> None:
    @app.before_request
    def complaint_security_guard():
        path = request.path
        complaint_path = path == "/complaint-audit" or path.startswith("/complaint-audit/") or path.startswith("/api/complaint-audit/")
        allowed_asset = path in {"/static/complaint_audit.css", "/static/complaint_audit.js"}
        if not is_loopback() and not complaint_path and not allowed_asset:
            return jsonify({"ok": False, "error": "内网远程访问仅开放已认证的客诉稽查看板"}), 403
        if not complaint_path or path in {"/complaint-audit/login", "/api/complaint-audit/auth/status"}:
            return None
        if not auth_configured():
            if path.startswith("/api/"):
                return jsonify({"ok": False, "error": "访问口令尚未配置"}), 503
            return redirect(url_for("complaint_audit.login"))
        if current_session() is None:
            return require_auth_response()
        if request.method not in {"GET", "HEAD", "OPTIONS"} and not validate_csrf():
            return jsonify({"ok": False, "error": "安全校验已失效，请刷新页面后重试"}), 403
        return None


@bp.route("/complaint-audit/login", methods=["GET", "POST"])
def login():
    error = ""
    if request.method == "POST":
        if not same_origin():
            error = "登录请求来源不受信任"
        elif not auth_configured():
            error = "访问口令尚未配置，请先运行本机配置程序"
        elif login_failure_count() >= LOGIN_FAILURE_LIMIT:
            append_access_log("限速")
            return render_template("complaint_audit_login.html", error="连续错误次数过多，请15分钟后再试", configured=True), 429
        else:
            password = request.form.get("password", "")
            try:
                verified = check_password_hash(os.environ["COMPLAINT_AUDIT_PASSWORD_HASH"], password)
            except (ValueError, KeyError):
                verified = False
            if verified:
                append_access_log("成功")
                response = make_response(redirect(url_for("complaint_audit.dashboard")))
                set_session_cookie(response)
                return response
            append_access_log("失败")
            error = "访问口令不正确"
    return render_template("complaint_audit_login.html", error=error, configured=auth_configured())


@bp.post("/complaint-audit/logout")
def logout():
    response = make_response(redirect(url_for("complaint_audit.login")))
    response.delete_cookie(SESSION_COOKIE, path="/")
    return response


@bp.get("/api/complaint-audit/auth/status")
def api_auth_status():
    return jsonify({"ok": True, "configured": auth_configured(), "authenticated": current_session() is not None})


@bp.get("/complaint-audit")
def dashboard():
    session = current_session() or {}
    return render_template(
        "complaint_audit.html", current_month=date.today().strftime("%Y-%m"),
        earliest_month=six_month_start().strftime("%Y-%m"), csrf_token=session.get("csrf", ""),
    )


@bp.get("/api/complaint-audit/summary")
def api_summary():
    try:
        filters = query_filters()
        all_records = load_aggregates(filters["month"])
        records = filter_records(all_records, filters)
        selected_start = month_start(filters["month"])
        previous_month = shift_month(selected_start, -1).strftime("%Y-%m")
        previous_filters = {**filters, "date_from": "", "date_to": "", "month": previous_month}
        previous = filter_records(load_aggregates(previous_month), previous_filters)
        coverage = month_coverage(filters["month"])
        previous_coverage = month_coverage(previous_month)
        granularity = "week" if request.args.get("granularity") == "week" else "day"
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    distribution, rankings = responsibility_analysis(records)
    roots = root_cause_analysis(records)
    not_upheld_causes = not_upheld_cause_analysis(records)
    not_upheld_descriptions = not_upheld_description_analysis(records)
    roots["quality"]["missing_completion_time"] = coverage["missing_completed_records"]
    counts = summary_counts(records)
    previous_counts = summary_counts(previous)
    counts["month_over_month"] = {
        "previous_month": previous_month, "inspected": previous_counts["inspected"],
        "upheld": previous_counts["upheld"], "not_upheld": previous_counts["not_upheld"],
        "inspected_delta": counts["inspected"] - previous_counts["inspected"],
        "comparable": coverage["complete_to_date"] and previous_coverage["complete_to_date"],
    }
    agents = sorted({agent for row in all_records for agent in row["complained_agents"]})
    brands = sorted({row["brand"] for row in all_records if row["brand"]})
    return jsonify({
        "ok": True,
        "data": {
            "month": filters["month"], "summary": counts,
            "trend": trend_rows(records, filters["month"], granularity), "granularity": granularity,
            "responsibility": distribution, "rankings": rankings[:50],
            "root_causes": roots,
            "not_upheld_causes": not_upheld_causes,
            "not_upheld_descriptions": not_upheld_descriptions,
            "facts": fact_summaries(records, previous, coverage, previous_coverage),
            "coverage": coverage,
            "options": {"agents": agents, "brands": brands, "responsibilities": list(RESPONSIBILITY_SUBJECTS)},
            "sync": sync_status(),
        },
    })


@bp.get("/api/complaint-audit/records")
def api_records():
    try:
        filters = query_filters()
        page = max(1, int(request.args.get("page", 1)))
        page_size = min(100, max(10, int(request.args.get("page_size", 25))))
        records = filter_records(load_aggregates(filters["month"]), filters)
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    records.sort(key=lambda row: (row["inspect_completed_at"], row["inspect_id"]), reverse=True)
    pages = max(1, math.ceil(len(records) / page_size))
    page = min(page, pages)
    offset = (page - 1) * page_size
    return jsonify({"ok": True, "data": {"total": len(records), "page": page, "pages": pages, "page_size": page_size, "records": records[offset:offset + page_size]}})


@bp.get("/api/complaint-audit/records/<inspect_id>/history")
def api_history(inspect_id: str):
    conn = get_db()
    rows = [dict(row) for row in conn.execute(
        "SELECT * FROM complaint_audit_versions WHERE inspect_id=? ORDER BY changed_at DESC,id DESC", (clean(inspect_id),)
    )]
    conn.close()
    for row in rows:
        fields = json.loads(row.pop("changed_fields"))
        before = json.loads(row.pop("before_json")) if row["before_json"] else None
        after = json.loads(row.pop("after_json"))
        row["changes"] = [
            {"field": field, "label": DISPLAY_NAMES.get(field, field),
             "before": before.get(field) if before else None, "after": after.get(field)}
            for field in fields
        ]
    return jsonify({"ok": True, "data": {"inspect_id": inspect_id, "versions": rows}})


@bp.get("/api/complaint-audit/sync/status")
def api_sync_status():
    session = current_session() or {}
    return jsonify({"ok": True, "data": sync_status(), "csrf_token": session.get("csrf", "")})


@bp.post("/api/complaint-audit/sync")
def api_sync():
    if source_mode() == "excel":
        return jsonify({"ok": False, "error": "当前为Excel导入模式，请上传最新Forge导出文件"}), 409
    if AUTO_STATE["running"]:
        return jsonify({"ok": False, "error": "已有客诉同步任务正在运行"}), 409
    payload = request.get_json(silent=True) or {}
    force_backfill = bool(payload.get("backfill"))

    def runner() -> None:
        try:
            sync_once(force_backfill=force_backfill)
        except Exception:
            pass

    threading.Thread(target=runner, daemon=True, name="complaint-audit-manual-sync").start()
    return jsonify({"ok": True, "message": "客诉增量同步已启动"})


@bp.post("/api/complaint-audit/import")
def api_import():
    upload = request.files.get("file")
    if not upload or not upload.filename:
        return jsonify({"ok": False, "error": "请选择Forge导出的Excel文件"}), 400
    if not upload.filename.lower().endswith(".xlsx"):
        return jsonify({"ok": False, "error": "仅支持.xlsx格式的Forge导出文件"}), 400
    content_length = request.content_length or 0
    if content_length > MAX_IMPORT_BYTES:
        return jsonify({"ok": False, "error": "导出文件超过25MB限制"}), 413
    try:
        result = import_excel_workbook(upload.stream, upload.filename)
    except (ValueError, KeyError) as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"ok": False, "error": f"导入失败：{exc}"}), 500
    return jsonify({"ok": True, "data": result})


def register_complaint_audit(app: Any) -> None:
    init_db()
    app.register_blueprint(bp)
    register_app_guards(app)
